import os
from docx import Document
from app.tools.resume_parser import extract_resume_text


def test_missing_file_returns_error():
    # There is no file at this path, so it should fail gracefully
    result = extract_resume_text("this_file_does_not_exist.pdf")
    assert result["success"] is False
    assert "not found" in result["error"].lower()


def test_unsupported_file_type_returns_error(tmp_path):
    # tmp_path is a temporary folder pytest creates and deletes automatically
    fake_file = tmp_path / "resume.txt"
    fake_file.write_text("Some resume content")

    result = extract_resume_text(str(fake_file))
    assert result["success"] is False
    assert "unsupported" in result["error"].lower()


def test_docx_extraction_returns_text(tmp_path):
    # Create a real, temporary .docx file to test against
    docx_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Email: jane.doe@example.com")
    document.add_paragraph("Skills: Python, SQL, Machine Learning")
    document.save(str(docx_path))

    result = extract_resume_text(str(docx_path))

    assert result["success"] is True
    assert "Jane Doe" in result["text"]
    assert "Python" in result["text"]
